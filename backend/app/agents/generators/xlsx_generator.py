"""Generates an actual working XLSX with real formulas (spec section 9)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def generate_budget_spreadsheet(path: Path, *, title: str, palette: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Budget"

    header_fill = PatternFill(start_color=palette.get("primary", "1B4332").lstrip("#"), fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=14)

    ws.merge_cells("A1:D1")
    ws["A1"] = title
    ws["A1"].font = header_font
    ws["A1"].fill = header_fill

    ws.append([])
    ws.append(["Category", "Planned", "Actual", "Difference"])
    for cell in ws[3]:
        cell.font = Font(bold=True)

    categories = ["Housing", "Groceries", "Transportation", "Utilities", "Entertainment", "Savings"]
    first_data_row = 4
    for i, cat in enumerate(categories):
        row = first_data_row + i
        ws.cell(row=row, column=1, value=cat)
        ws.cell(row=row, column=2, value=0)
        ws.cell(row=row, column=3, value=0)
        # REAL working formula: Difference = Planned - Actual
        ws.cell(row=row, column=4, value=f"=B{row}-C{row}")

    total_row = first_data_row + len(categories)
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=total_row, column=2, value=f"=SUM(B{first_data_row}:B{total_row - 1})")
    ws.cell(row=total_row, column=3, value=f"=SUM(C{first_data_row}:C{total_row - 1})")
    ws.cell(row=total_row, column=4, value=f"=SUM(D{first_data_row}:D{total_row - 1})")

    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = 20

    wb.save(path)
